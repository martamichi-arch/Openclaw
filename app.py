import io
import os
import tempfile
from pathlib import Path
from datetime import datetime

import streamlit as st
import fitz
from pptx import Presentation
from docx import Document
from docx.shared import Pt, Inches

try:
    from openai import OpenAI
except Exception:
    OpenAI = None


st.set_page_config(
    page_title="StudyBook Builder",
    page_icon="📚",
    layout="wide"
)

st.title("📚 StudyBook Builder")
st.write("Carica appunti, slide e testi. Genera un documento Word unico.")


def get_secret(name):
    try:
        return st.secrets.get(name, "")
    except Exception:
        return os.getenv(name, "")


def clean(text):
    if not text:
        return ""
    lines = []
    for line in text.splitlines():
        line = line.strip()
        if line:
            lines.append(line)
    return "\n".join(lines)


def pdf_extract(file, category):
    data = file.read()
    pdf = fitz.open(stream=data, filetype="pdf")

    item = {
        "cat": category,
        "name": file.name,
        "texts": [],
        "imgs": []
    }

    for page_index, page in enumerate(pdf):
        page_number = page_index + 1
        source = file.name + " - pagina " + str(page_number)
        text = clean(page.get_text("text"))

        if text:
            item["texts"].append((source, text))

        for image_index, image_info in enumerate(page.get_images(full=True)):
            try:
                xref = image_info[0]
                base = pdf.extract_image(xref)
                ext = base.get("ext", "png")

                filename = (
                    Path(file.name).stem
                    + "_p"
                    + str(page_number)
                    + "_"
                    + str(image_index + 1)
                    + "."
                    + ext
                )

                item["imgs"].append(
                    (source, filename, base["image"], text[:800])
                )
            except Exception:
                pass

    return item


def docx_extract(file, category):
    data = file.read()
    document = Document(io.BytesIO(data))

    item = {
        "cat": category,
        "name": file.name,
        "texts": [],
        "imgs": []
    }

    paragraphs = []
    for paragraph in document.paragraphs:
        text = clean(paragraph.text)
        if text:
            paragraphs.append(text)

    full_text = "\n".join(paragraphs)

    if full_text:
        item["texts"].append((file.name, full_text))

    image_count = 0

    for rel_key in document.part._rels:
        try:
            rel = document.part._rels[rel_key]
            if "image" in rel.target_ref:
                image_count += 1
                filename = (
                    Path(file.name).stem
                    + "_img"
                    + str(image_count)
                    + ".png"
                )
                item["imgs"].append(
                    (file.name, filename, rel.target_part.blob, full_text[:800])
                )
        except Exception:
            pass

    return item


def pptx_extract(file, category):
    data = file.read()
    presentation = Presentation(io.BytesIO(data))

    item = {
        "cat": category,
        "name": file.name,
        "texts": [],
        "imgs": []
    }

    for slide_index, slide in enumerate(presentation.slides):
        slide_number = slide_index + 1
        source = file.name + " - slide " + str(slide_number)

        text_parts = []

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = clean(shape.text)
                if text:
                    text_parts.append(text)

        slide_text = "\n".join(text_parts)

        if slide_text:
            item["texts"].append((source, slide_text))

        image_count = 0

        for shape in slide.shapes:
            try:
                if hasattr(shape, "image"):
                    image_count += 1
                    image = shape.image
                    ext = image.ext or "png"

                    filename = (
                        Path(file.name).stem
                        + "_slide"
                        + str(slide_number)
                        + "_"
                        + str(image_count)
                        + "."
                        + ext
                    )

                    item["imgs"].append(
                        (source, filename, image.blob, slide_text[:800])
                    )
            except Exception:
                pass

    return item


def txt_extract(file, category):
    raw = file.read()

    try:
        text = raw.decode("utf-8")
    except Exception:
        text = raw.decode("latin-1", errors="ignore")

    text = clean(text)

    item = {
        "cat": category,
        "name": file.name,
        "texts": [],
        "imgs": []
    }

    if text:
        item["texts"].append((file.name, text))

    return item


def extract_file(file, category):
    extension = Path(file.name).suffix.lower()

    if extension == ".pdf":
        return pdf_extract(file, category)

    if extension == ".docx":
        return docx_extract(file, category)

    if extension == ".pptx":
        return pptx_extract(file, category)

    if extension == ".txt":
        return txt_extract(file, category)

    return {
        "cat": category,
        "name": file.name,
        "texts": [],
        "imgs": []
    }


def split_text(text, max_length=9000):
    chunks = []
    current = ""

    for paragraph in text.split("\n"):
        if len(current) + len(paragraph) > max_length and current:
            chunks.append(current)
            current = ""

        current += paragraph + "\n"

    if current:
        chunks.append(current)

    return chunks


def ai_rewrite(blocks, notes):
    api_key = get_secret("OPENAI_API_KEY")

    if not api_key or OpenAI is None:
        return None

    client = OpenAI(api_key=api_key)

    raw_parts = []

    for source, text in blocks:
        raw_parts.append("[FONTE: " + source + "]\n" + text)

    raw_text = "\n\n".join(raw_parts)
    pieces = split_text(raw_text)

    outputs = []

    for index, piece in enumerate(pieces):
        prompt = "\n".join([
            "Trasforma questi materiali in un testo unico di studio in italiano.",
            "Regole obbligatorie:",
            "- non omettere informazioni",
            "- non inventare",
            "- non generare immagini",
            "- usa prima gli appunti studenti",
            "- integra poi le slide del professore",
            "- usa infine i testi di riferimento per chiarire",
            "- scrivi come un libro di testo",
            "- organizza in capitoli e paragrafi",
            "Blocco " + str(index + 1) + " di " + str(len(pieces)),
            "Rettifiche utente:",
            notes or "Nessuna.",
            "Materiale:",
            piece
        ])

        response = client.chat.completions.create(
            model="gpt-4.1-mini",
            messages=[
                {
                    "role": "system",
                    "content": "Ricostruisci dispense universitarie fedeli, complete e chiare."},
                {
                    "role": "user",
                    "content": prompt}],
            temperature=0.2)

        outputs.append(response.choices[0].message.content)

    return "\n\n".join(outputs)


def add_paragraph(document, text, size=10, bold=False):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.bold = bold
    return paragraph


def add_heading(document, text, level=1):
    if level == 1:
        size = 16
    elif level == 2:
        size = 13
    else:
        size = 11

    add_paragraph(document, text, size=size, bold=True)


def make_docx(items, ai_text, notes):
    document = Document()

    document.styles["Normal"].font.name = "Aptos"
    document.styles["Normal"].font.size = Pt(10)

    add_heading(document, "Documento unico di studio", 1)
    add_paragraph(
        document,
        "Generato il " + datetime.now().strftime("%d/%m/%Y %H:%M"),
        size=9
    )

    if notes:
        add_heading(document, "Rettifiche richieste", 2)
        add_paragraph(document, notes)

    document.add_page_break()
    add_heading(document, "Testo ricostruito", 1)

    if ai_text:
        for line in ai_text.splitlines():
            line = line.strip()
            if line:
                add_paragraph(document, line)
    else:
        add_paragraph(
            document,
            "AI non attiva: documento estrattivo.",
            bold=True
        )

        for item in items:
            add_heading(document, item["cat"] + " - " + item["name"], 2)

            for source, text in item["texts"]:
                add_heading(document, source, 3)
                add_paragraph(document, text)

    document.add_page_break()
    add_heading(document, "Immagini originali estratte dai file", 1)
    add_paragraph(
        document,
        "Nessuna immagine è stata generata da AI.",
        size=9
    )

    temporary_directory = tempfile.mkdtemp()

    for item in items:
        for source, filename, image_bytes, context in item["imgs"]:
            try:
                image_path = Path(temporary_directory) / filename
                image_path.write_bytes(image_bytes)

                add_heading(document, source, 3)
                document.add_picture(str(image_path), width=Inches(5.8))

                if context:
                    add_paragraph(
                        document,
                        "Contesto: " + context,
                        size=9
                    )
            except Exception:
                pass

    document.add_page_break()
    add_heading(document, "Appendice - testo estratto integralmente", 1)

    for item in items:
        add_heading(document, item["cat"] + " - " + item["name"], 2)

        for source, text in item["texts"]:
            add_heading(document, source, 3)
            add_paragraph(document, text, size=9)

    output = io.BytesIO()
    document.save(output)
    output.seek(0)

    return output


with st.sidebar:
    st.header("Stato")

    if get_secret("OPENAI_API_KEY") and OpenAI is not None:
        st.success("AI attiva")
    else:
        st.warning("AI non attiva: documento estrattivo")


student_files = st.file_uploader(
    "APPUNTI STUDENTI",
    type=["pdf", "docx", "pptx", "txt"],
    accept_multiple_files=True
)

slide_files = st.file_uploader(
    "SLIDE PROFESSORE",
    type=["pdf", "docx", "pptx", "txt"],
    accept_multiple_files=True
)

reference_files = st.file_uploader(
    "TESTI DI RIFERIMENTO",
    type=["pdf", "docx", "pptx", "txt"],
    accept_multiple_files=True
)

extra_files = st.file_uploader(
    "APPUNTI / EXTRA",
    type=["pdf", "docx", "pptx", "txt"],
    accept_multiple_files=True
)

notes = st.text_area(
    "Rettifiche o istruzioni aggiuntive",
    height=120
)

if st.button("Genera documento Word", type="primary"):
    groups = [
        ("APPUNTI STUDENTI", student_files),
        ("SLIDE PROFESSORE", slide_files),
        ("TESTI DI RIFERIMENTO", reference_files),
        ("APPUNTI / EXTRA", extra_files)
    ]

    items = []

    with st.spinner("Estraggo testo e immagini..."):
        for category, files in groups:
            for file in files or []:
                try:
                    items.append(extract_file(file, category))
                except Exception as error:
                    st.error("Errore in " + file.name + ": " + str(error))

    blocks = []

    for item in items:
        for source, text in item["texts"]:
            blocks.append((item["cat"] + " - " + source, text))

    if not blocks:
        st.error("Non ho trovato testo nei file caricati.")
        st.stop()

    with st.spinner("Ricostruisco il testo..."):
        ai_text = ai_rewrite(blocks, notes)

    with st.spinner("Creo il Word..."):
        result = make_docx(items, ai_text, notes)

    st.success("Documento generato.")

    st.download_button(
        "Scarica documento Word",
        result,
        "documento_unico_studio.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
